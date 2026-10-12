"""reference.docx ni qayta yasaydi: python3 kitob/shablon/reference_yasash.py

Pandoc'ning standart reference.docx faylidan boshlab shriftlar, A4 va bet raqamini sozlaydi.
Uslublarni o'zgartirish uchun shu faylni tahrirlab qayta ishga tushiring.
"""
import copy
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CHIQISH = Path(__file__).with_name("reference.docx")
SERIF = "Times New Roman"
SANS = "Arial"


def shrift(stil, nom, hajm=None, qalin=None, rang=None, kursiv=None):
    f = stil.font
    f.name = nom
    rpr = stil.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), nom)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(a), None)
    if hajm:
        f.size = Pt(hajm)
    if qalin is not None:
        f.bold = qalin
    if kursiv is not None:
        f.italic = kursiv
    if rang:
        f.color.rgb = RGBColor.from_string(rang)


def bet_oldin(stil):
    ppr = stil.element.get_or_add_pPr()
    ppr.append(OxmlElement("w:pageBreakBefore"))


def chap_chiziq(stil, rang):
    ppr = stil.element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "18"), ("space", "8"), ("color", rang)):
        left.set(qn(f"w:{k}"), v)
    pbdr.append(left)
    ppr.append(pbdr)


def bet_raqami(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tur, matn in (("begin", None), (None, " PAGE "), ("separate", None), ("end", None)):
        if tur:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tur)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = matn
        run._r.append(el)
        if tur == "separate":
            t = OxmlElement("w:t")
            t.text = "1"
            run._r.append(t)
    run.font.name = SANS
    run.font.size = Pt(9)


def main():
    with tempfile.TemporaryDirectory() as t:
        boshlang = Path(t) / "ref.docx"
        with open(boshlang, "wb") as f:
            f.write(subprocess.run(
                ["pandoc", "--print-default-data-file", "reference.docx"],
                capture_output=True, check=True).stdout)
        d = Document(str(boshlang))

    s = {st.name: st for st in d.styles}
    stillar = d.styles
    for nom in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text", "Definition"):
        shrift(s[nom], SERIF, 11)
    for nom in ("Body Text", "First Paragraph"):
        pf = s[nom].paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(6), 1.15
    s["Compact"].paragraph_format.space_after = Pt(2)
    s["Normal"].paragraph_format.line_spacing = 1.15

    for nom, hajm, oldin, keyin in (("Heading 1", 22, 0, 18), ("Heading 2", 15, 16, 6),
                                    ("Heading 3", 12, 12, 4), ("Heading 4", 11, 10, 3)):
        shrift(s[nom], SANS, hajm, qalin=True, rang="1F3A5F")
        s[nom].paragraph_format.space_before = Pt(oldin)
        s[nom].paragraph_format.space_after = Pt(keyin)
        s[nom].paragraph_format.keep_with_next = True

    shrift(s["Title"], SANS, 34, qalin=True, rang="1F3A5F")
    s["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s["Title"].paragraph_format.space_before = Pt(150)
    s["Title"].paragraph_format.space_after = Pt(12)
    shrift(s["Subtitle"], SANS, 18, rang="444444")
    s["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s["Subtitle"].paragraph_format.space_after = Pt(120)
    shrift(s["Author"], SANS, 16, rang="222222")
    s["Author"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shrift(s["Date"], SANS, 12, rang="444444")
    s["Date"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    shrift(s["TOC Heading"], SANS, 22, qalin=True, rang="1F3A5F")
    bet_oldin(s["TOC Heading"])  # mundarija muqovadan keyin yangi betda
    shrift(s["Caption"], SANS, 9, rang="444444")
    shrift(s["Image Caption"], SANS, 9, rang="444444")
    shrift(s["Table Caption"], SANS, 9, rang="444444")

    # Filtr beradigan uslublar
    for nom, chap, indent in (("Mashq", None, 0), ("SAT savol", "B8860B", 0.4), ("Misol", "1F6FEB", 0.4)):
        st = s[nom] = stillar.add_style(nom, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = s["Body Text"]
        st.paragraph_format.space_after = Pt(6)
        if indent:
            st.paragraph_format.left_indent = Cm(indent)
        if chap:
            chap_chiziq(st, chap)
    s["Mashq"].paragraph_format.keep_together = True
    s["SAT savol"].paragraph_format.keep_with_next = True

    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    bet_raqami(sec)
    d.save(str(CHIQISH))


if __name__ == "__main__":
    main()
