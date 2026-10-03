# General English kurs tuzilmasi — SAT_Math_kurs_tuzilmasi.xlsx bilan bir xil tuzilma va uslubda
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = 'General_English_kurs_tuzilmasi.xlsx'
COURSE = 'General English: A2 → B1'
INDIGO = 'FF4F46E5'
thin = Side(style='thin', color='FFBFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FONT = Font(name='Arial', size=11, bold=True, color='FFFFFFFF')
HEAD_FILL = PatternFill('solid', fgColor=INDIGO)
YELLOW = PatternFill('solid', fgColor='FFFFF9C4')
DOMAIN_FILL = {
    'Kirish / Takrorlash': 'FFEEEEEE',
    'Grammar': 'FFE8F1FA',
    'Vocabulary': 'FFEDE7F6',
    'Reading & Listening': 'FFE6F4EA',
    'Writing & Speaking': 'FFFFF4E0',
}

MODULES = [
    ('M0', 'Kirish: darajani aniqlash va o\'rganish usuli', '—', 'Diagnostik test'),
    ('M1', 'People & Daily Life (odamlar va kundalik hayot)', 'A2', 'Modul testi (30 savol)'),
    ('M2', 'Past & Experiences (o\'tmish va tajriba)', 'A2+', 'Modul testi (30 savol)'),
    ('M3', 'Future, Plans & Work (kelajak, rejalar, ish)', 'B1−', 'Modul testi (30 savol)'),
    ('M4', 'Rules, Advice & Conditions (qoidalar, maslahat, shartlar)', 'B1', 'Modul testi (30 savol)'),
    ('M5', 'Takrorlash va yakuniy testlar', 'B1', 'Mock testlar'),
]

T10 = '10 savol, 15 min'
TW = '10 savol + yozma ish (80–120 so\'z)'
TS = '10 savol + 1 daqiqalik audio javob'
MOCK = 'Mock: 60 savol, 70 min + Writing'
PR = 'Xatolar ustida ishlash'
PRW = 'Yozma ish (o\'qituvchi tekshiradi)'
PRS = 'Audio javob (o\'qituvchi tekshiradi)'

# (modul, nomi, savollar bazasidagi mavzu, yo'nalish, inglizcha kalit so'zlar, nima o'rgatiladi, CEFR, mashq, test)
L = [
    ('M0', 'Kursga kirish va darajani aniqlash testi', 'English — Diagnostika', 'Kirish / Takrorlash',
     'level, placement test, CEFR, skill, grammar, vocabulary',
     "CEFR darajalari (A1–C2) va 4 ko'nikma; diagnostik test; natijadan shaxsiy o'rganish rejasi", 'A2–B1', PR, 'Diagnostik: 40 savol, 30 min'),
    ('M0', "Qanday o'rganamiz: lug'at daftari, talaffuz va takrorlash", 'English — Diagnostika', 'Kirish / Takrorlash',
     'vocabulary notebook, collocation, word stress, phonetic symbols, review',
     "So'zni kollokatsiya va misol bilan yozish, IPA belgilari asoslari, interval takrorlash tizimi", 'A2', PR, T10),
    ('M1', 'Present Simple: kundalik hayot va odatlar', 'Present Simple', 'Grammar',
     'always, usually, often, sometimes, never, every day, does / doesn\'t',
     "Tasdiq, inkor, so'roq; 3-shaxs -s; chastota ravishlarining gapdagi o'rni", 'A2', PR, T10),
    ('M1', 'Present Continuous va Present Simple farqi', 'Present Continuous', 'Grammar',
     'now, at the moment, currently, these days, state verbs (like, know, want)',
     "Hozirgi jarayon va odat; holat fe'llari Continuous'da ishlatilmasligi", 'A2', PR, T10),
    ('M1', "Lug'at: oila, tashqi ko'rinish va xarakter", 'Vocabulary — People and family', 'Vocabulary',
     'relatives, tall, curly hair, friendly, generous, look like, be like',
     "Odamni tasvirlash; «What does she look like?» va «What is she like?» farqi", 'A2', PR, T10),
    ('M1', "Savol tuzish: so'roq so'zlar va yordamchi fe'llar", 'Questions and question words', 'Grammar',
     'who, what, where, how often, how long, subject questions',
     "Savol so'z tartibi (QASI), subject va object savollar, qisqa javoblar", 'A2', PR, T10),
    ('M1', "Otlar, artikllar va miqdor so'zlari", 'Articles and quantifiers', 'Grammar',
     'a / an, the, zero article, some, any, much, many, a lot of, countable',
     "Sanaladigan / sanalmaydigan otlar; a/an/the qoidalari; some/any, much/many", 'A2', PR, T10),
    ('M1', 'Reading: qisqa matndan asosiy fikr va detallarni topish', 'Reading — Main idea and details', 'Reading & Listening',
     'skim, scan, main idea, detail, true / false / not given',
     "Skimming va scanning; asosiy fikr va detalni farqlash; True / False / Not Given", 'A2', PR, T10),
    ('M1', "Speaking: o'zini tanishtirish va small talk", 'Functional language — Introductions', 'Writing & Speaking',
     'Nice to meet you, What do you do?, How about you?, follow-up question',
     "Tanishuv iboralari, suhbatni davom ettirish uchun qo'shimcha savollar, intonatsiya", 'A2', PRS, TS),
    ('M1', "Writing: do'stga email + M1 modul testi", 'Writing — Informal email', 'Writing & Speaking',
     'Hi …, Thanks for …, Guess what!, Hope to hear from you, and / but / because / so',
     "Norasmiy email tuzilmasi va bog'lovchilar; modul testi", 'A2', PRW, TW),
    ('M2', "Past Simple: to'g'ri va noto'g'ri fe'llar", 'Past Simple', 'Grammar',
     'yesterday, last week, ago, -ed, went, saw, did / didn\'t',
     "-ed talaffuzi (/t/ /d/ /ɪd/), 50 ta asosiy noto'g'ri fe'l, inkor va so'roq", 'A2', PR, T10),
    ('M2', 'Past Continuous va Past Simple: voqeani hikoya qilish', 'Past Continuous', 'Grammar',
     'while, when, was / were + -ing, suddenly, at that moment',
     "Fon jarayoni va to'satdan yuz bergan voqea; when / while bilan gaplar", 'A2+', PR, T10),
    ('M2', "Used to va would: o'tmishdagi odatlar", 'Used to', 'Grammar',
     'used to, didn\'t use to, would, any more, no longer',
     "O'tmishdagi odat va holat; used to va would farqi; be used to bilan adashtirmaslik", 'A2+', PR, T10),
    ('M2', 'Present Perfect: tajriba (ever, never, just, already, yet)', 'Present Perfect', 'Grammar',
     'ever, never, just, already, yet, been / gone',
     "Hayotiy tajriba va yangi natija; III shakl; been va gone farqi", 'A2+', PR, T10),
    ('M2', 'Present Perfect yoki Past Simple? (for, since)', 'Present Perfect vs Past Simple', 'Grammar',
     'for, since, how long, last year, ago, so far',
     "Aniq o'tgan vaqt — Past Simple; hozirgacha davom — Present Perfect; for va since", 'A2+', PR, T10),
    ('M2', "Lug'at: sayohat, transport va dam olish", 'Vocabulary — Travel and holidays', 'Vocabulary',
     'book a ticket, check in, journey / trip / travel, sightseeing, delay, souvenir',
     "Sayohat kollokatsiyalari; journey, trip, travel farqi; aeroport va mehmonxona iboralari", 'A2+', PR, T10),
    ('M2', "Listening: e'lon va suhbatdan aniq ma'lumot olish", 'Listening — Specific information', 'Reading & Listening',
     'predict, key words, numbers and dates, spelling, distractor',
     "Tinglashdan oldin bashorat qilish, raqam va sanalarni eshitish, chalg'ituvchi javoblar", 'A2+', PR, T10),
    ('M2', 'Writing: hikoya yozish + M2 modul testi', 'Writing — Narrative', 'Writing & Speaking',
     'first, then, after that, suddenly, finally, in the end',
     "Hikoya tuzilmasi (boshlanish, voqea, yakun), vaqt bog'lovchilari, o'tgan zamonlar; modul testi", 'A2+', PRW, TW),
    ('M3', 'Kelasi zamon: will, be going to va Present Continuous', 'Future forms', 'Grammar',
     'will, be going to, I\'m meeting …, prediction, plan, arrangement',
     "Bashorat, reja va kelishuvni farqlash; to'satdan qaror (will)", 'B1−', PR, T10),
    ('M3', 'Sifat darajalari: comparatives va superlatives', 'Comparatives and superlatives', 'Grammar',
     '-er than, more … than, the -est, as … as, much / a bit, not as … as',
     "Qiyosiy va orttirma daraja, noto'g'ri shakllar (good / better / best), as … as", 'B1−', PR, T10),
    ('M3', "Modal fe'llar: can, could, be able to", 'Modals of ability', 'Grammar',
     'can, could, be able to, manage to, will be able to',
     "Hozirgi, o'tgan va kelasi qobiliyat; could va was able to farqi", 'B1−', PR, T10),
    ('M3', "Lug'at: ish, kasb va ta'lim", 'Vocabulary — Work and study', 'Vocabulary',
     'apply for, job interview, salary, career, degree, take / pass an exam',
     "Ish va o'qishga oid kollokatsiyalar; make / do / take bilan iboralar", 'B1−', PR, T10),
    ('M3', "Fe'l + -ing yoki to-infinitive", 'Gerunds and infinitives', 'Grammar',
     'enjoy -ing, want to, decide to, avoid -ing, stop to / stop -ing',
     "Qaysi fe'ldan keyin -ing, qaysidan keyin to + fe'l; ma'no o'zgaradigan fe'llar", 'B1', PR, T10),
    ('M3', "Reading: muallif fikri va so'z ma'nosini kontekstdan topish", 'Reading — Inference and vocabulary in context', 'Reading & Listening',
     'infer, imply, opinion, fact, synonym, context clue',
     "Matnda yashirin fikrni topish, fakt va fikrni ajratish, notanish so'zni kontekstdan aniqlash", 'B1', PR, T10),
    ('M3', 'Speaking: rejalar, taklif va kelishish', 'Functional language — Suggestions and arrangements', 'Writing & Speaking',
     'Shall we …?, How about …?, Why don\'t we …?, I\'d rather …, That sounds good',
     "Taklif berish, rozilik va muloyim rad etish, uchrashuv vaqtini kelishish", 'B1', PRS, TS),
    ('M3', 'Writing: rasmiy email va ariza + M3 modul testi', 'Writing — Formal email', 'Writing & Speaking',
     'Dear Sir or Madam, I am writing to enquire, I would be grateful, Yours faithfully',
     "Rasmiy va norasmiy uslub farqi, ish yoki kursga ariza; modul testi", 'B1', PRW, TW),
    ('M4', 'Majburiyat va maslahat: must, have to, should', 'Modals of obligation and advice', 'Grammar',
     'must, have to, don\'t have to, mustn\'t, should, ought to',
     "Majburiyat, taqiq va zarurat yo'qligi (mustn't ≠ don't have to); maslahat berish", 'B1', PR, T10),
    ('M4', 'Zero va First Conditional', 'Conditionals (zero and first)', 'Grammar',
     'if, unless, when, as soon as, will, real situation',
     "Umumiy haqiqat va real kelajak sharti; if-gapda will ishlatilmasligi; unless", 'B1', PR, T10),
    ('M4', 'Second Conditional: xayoliy vaziyatlar', 'Second conditional', 'Grammar',
     'If I were you, would, could, imaginary, wish',
     "Xayoliy hozirgi/kelajak; first va second conditional farqi; maslahat: If I were you …", 'B1', PR, T10),
    ('M4', 'Passive Voice: Present va Past', 'Passive voice', 'Grammar',
     'is made, was built, by, active, passive, was invented',
     "Majhul nisbat qachon ishlatiladi; is/was + III shakl; by bilan bajaruvchi", 'B1', PR, T10),
    ('M4', "Lug'at: kundalik phrasal verbs", 'Vocabulary — Phrasal verbs', 'Vocabulary',
     'look for, give up, find out, turn on / off, get up, put off',
     "30 ta eng ko'p ishlatiladigan phrasal verb; ajraladigan va ajralmaydiganlar", 'B1', PR, T10),
    ('M4', 'Relative clauses: who, which, that, where, whose', 'Relative clauses', 'Grammar',
     'who, which, that, where, whose, defining clause',
     "Odam, narsa va joyni aniqlovchi ergash gaplar; that qachon tushib qoladi", 'B1', PR, T10),
    ('M4', 'Speaking: fikr bildirish va asoslash + M4 modul testi', 'Functional language — Giving opinions', 'Writing & Speaking',
     'In my opinion, I agree, I see your point, but …, because, for example',
     "Fikr bildirish, rozi bo'lish/bo'lmaslik, misol bilan asoslash; modul testi", 'B1', PRS, TS),
    ('M5', 'Reported speech va zamonlar xaritasi + Mock 1', 'English — Mock testlar', 'Kirish / Takrorlash',
     'said, told, asked, reported speech, tense review, timeline',
     "Ko'chirma gap asoslari (said / told), 6 ta zamonni bitta xaritada takrorlash; to'liq mock test", 'B1', PR, MOCK),
    ('M5', 'Imtihon strategiyasi: Reading va Listening + Mock 2', 'English — Mock testlar', 'Kirish / Takrorlash',
     'time management, skim, key words, eliminate, guess, check',
     "Vaqtni taqsimlash, variantlarni chiqarib tashlash, javoblarni tekshirish; to'liq mock test", 'B1', PR, MOCK),
    ('M5', 'Yakuniy Mock 3, Writing/Speaking baholash va natija tahlili', 'English — Mock testlar', 'Kirish / Takrorlash',
     'score report, band, feedback, accuracy, fluency, progress',
     "Yakuniy mock, Writing va Speaking baholash mezonlari, diagnostika bilan solishtirish, sertifikat", 'B1', PR, MOCK),
]
assert len(L) == 36

wb = Workbook()

def header(ws, cols, height=26.85):
    for i, h in enumerate(cols, 1):
        c = ws.cell(1, i, h)
        c.font = HEAD_FONT; c.fill = HEAD_FILL; c.border = BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = height

def widths(ws, w):
    for k, v in w.items():
        ws.column_dimensions[k].width = v

N = len(L) + 1  # oxirgi qator (Darslar)
RNG = lambda col: f"Darslar!${col}$2:${col}${N}"

# ---------- Kurs kartochkasi ----------
ws = wb.active
ws.title = 'Kurs kartochkasi'
ws['A1'] = f'{COURSE} — kurs kartochkasi'
ws['A1'].font = Font(name='Arial', size=14, bold=True, color=INDIGO)
card = [
    ('Kurs nomi', COURSE),
    ('Qisqa tavsif (sayt uchun)', "Umumiy ingliz tili: A2 (Elementary) darajasidan B1 (Intermediate) darajasigacha. 6 ta modul, 36 ta dars, 3 ta yakuniy mock test. Grammatika o'zbek tilida tushuntiriladi, har darsda lug'at, talaffuz va amaliy ko'nikma (reading, listening, writing yoki speaking)."),
    ('Davomiyligi', '12 hafta, haftasiga 3 ta dars (36 dars)'),
    ('Kimlar uchun', "Ingliz tilini asosiy darajada biladigan (A2) o'quvchilar, maktab o'quvchilari va kattalar; IELTS / CEFR / SAT ga tayyorlanishdan oldingi bosqich"),
    ('Natija', "Kirish va chiqish testlari solishtiriladi; kurs 100% tugatilganda sertifikat. Eslatma: B1 ga to'liq yetish uchun darslardan tashqari muntazam mashq (kuniga 20–30 daqiqa) kerak."),
    ('Daraja va format', "CEFR A2 → B1. Har bir dars: video (10–15 min) + konspekt + dars testi. Modul oxirida 30 savollik modul testi. Mock: 60 savol (Grammar & Vocabulary 30, Reading 15, Listening 15), 70 min + Writing (120–150 so'z)."),
    ('Yo\'nalishlar ulushi', None),
    ('Manba', "Council of Europe — CEFR Companion Volume (2020), coe.int/lang-cefr; grammatika tartibi: British Council & EAQUALS, Core Inventory for General English (2010)"),
]
for i, (k, v) in enumerate(card, start=3):
    ws.cell(i, 1, k).font = Font(name='Arial', size=11, bold=True)
    c = ws.cell(i, 2, v); c.font = Font(name='Arial', size=11)
    c.alignment = Alignment(wrap_text=True, vertical='top'); ws.cell(i, 1).alignment = Alignment(vertical='top')
doms = list(DOMAIN_FILL)
parts = [f'"{d} "&COUNTIF({RNG("F")},"{d}")&" dars"' for d in ['Grammar', 'Vocabulary', 'Reading & Listening', 'Writing & Speaking', 'Kirish / Takrorlash']]
ws['B9'] = '=' + '&" · "&'.join(parts)
ws['B9'].comment = Comment("Avtomatik: «Darslar» varag'idagi «Yo'nalish» ustunidan hisoblanadi.", 'Kholmurodov Academy')
for r, h in {4: 48, 7: 33, 8: 48, 9: 33, 10: 33}.items():
    ws.row_dimensions[r].height = h
widths(ws, {'A': 24, 'B': 90})

# ---------- Modullar ----------
ws = wb.create_sheet('Modullar')
header(ws, ['Modul', 'Modul nomi', 'CEFR daraja', 'Darslar soni', 'Haftalar', 'Modul oxirida'])
for r, (m, name, lvl, end) in enumerate(MODULES, start=2):
    mn = f'_xlfn.MINIFS({RNG("C")},{RNG("B")},A{r})'
    mx = f'_xlfn.MAXIFS({RNG("C")},{RNG("B")},A{r})'
    vals = [m, name, lvl, f'=COUNTIF({RNG("B")},A{r})', f'=IF({mn}={mx},{mn}&"",{mn}&"–"&{mx})', end]
    for c, v in enumerate(vals, 1):
        x = ws.cell(r, c, v); x.font = Font(name='Arial', size=11); x.border = BORDER
        x.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center' if c in (1, 3, 4, 5) else None)
r = len(MODULES) + 2
ws.cell(r, 1, 'Jami').font = Font(name='Arial', size=11, bold=True)
ws.cell(r, 4, f'=SUM(D2:D{r-1})').font = Font(name='Arial', size=11, bold=True)
for c in range(1, 7):
    ws.cell(r, c).border = BORDER
ws.cell(r, 4).alignment = Alignment(horizontal='center')
widths(ws, {'A': 8, 'B': 52, 'C': 14, 'D': 12, 'E': 10, 'F': 24})

# ---------- Darslar ----------
ws = wb.create_sheet('Darslar')
header(ws, ['Dars №', 'Modul', 'Hafta', 'Dars nomi (sayt uchun)', 'Savollar bazasidagi mavzu', "Yo'nalish",
            "Inglizcha kalit so'zlar", "Darsda nima o'rgatiladi", 'CEFR', 'Video', 'Mashq (platforma)', 'Dars testi'], 39.55)
for i, (m, name, topic, dom, kw, what, lvl, pr, test) in enumerate(L, start=1):
    r = i + 1
    vals = [i, m, math.ceil(i / 3), name, topic, dom, kw, what, lvl, '10–15 min', pr, test]
    for c, v in enumerate(vals, 1):
        x = ws.cell(r, c, v); x.font = Font(name='Arial', size=10); x.border = BORDER
        x.alignment = Alignment(vertical='top', wrap_text=c in (4, 5, 6, 7, 8, 11, 12),
                                horizontal='center' if c in (1, 2, 3, 9) else None)
    ws.cell(r, 6).fill = PatternFill('solid', fgColor=DOMAIN_FILL[dom])
    ws.row_dimensions[r].height = 38
ws.freeze_panes = 'E2'
ws.auto_filter.ref = f'A1:L{N}'
widths(ws, {'A': 7, 'B': 7, 'C': 7, 'D': 38, 'E': 34, 'F': 20, 'G': 36, 'H': 44, 'I': 9, 'J': 10, 'K': 22, 'L': 22})

# ---------- Savollar bazasi ----------
ws = wb.create_sheet('Savollar bazasi')
header(ws, ['Mavzu (platformadagi nom — aynan shunday)', "Yo'nalish", 'Darslar soni', 'Maqsad: savollar', 'Yozilgan savollar', 'Holat'])
topics = []
for row in L:
    if (row[2], row[3]) not in topics:
        topics.append((row[2], row[3]))
TARGET = {'Grammar': 60, 'Vocabulary': 50, 'Reading & Listening': 40, 'Writing & Speaking': 30}
for r, (t, d) in enumerate(topics, start=2):
    target = 40 if 'Diagnostika' in t else 180 if 'Mock' in t else TARGET[d]
    vals = [t, d, f'=COUNTIF({RNG("E")},A{r})', target, None, f'=IF(E{r}="","Boshlanmagan",IF(E{r}>=D{r},"Tayyor","Jarayonda"))']
    for c, v in enumerate(vals, 1):
        x = ws.cell(r, c, v); x.font = Font(name='Arial', size=10); x.border = BORDER
        x.alignment = Alignment(wrap_text=(c == 1), vertical='top', horizontal='center' if c in (3, 4, 5) else None)
    ws.cell(r, 2).fill = PatternFill('solid', fgColor=DOMAIN_FILL[d])
    ws.cell(r, 5).fill = YELLOW
last = len(topics) + 1
r = last + 1
for c, v in enumerate(['Jami', None, f'=SUM(C2:C{last})', f'=SUM(D2:D{last})', f'=SUM(E2:E{last})', None], 1):
    x = ws.cell(r, c, v); x.font = Font(name='Arial', size=10, bold=True); x.border = BORDER
    x.alignment = Alignment(horizontal='center' if c in (3, 4, 5) else None)
note = ws.cell(r + 1, 1, "Sariq katakka yozilgan savollar sonini kiriting — «Holat» o'zi yangilanadi. Maqsad: Grammar mavzulariga 60 tadan, Vocabulary — 50, Reading & Listening — 40, Writing & Speaking — 30 (test qismi); diagnostika — 40; mock uchun 3 ta to'liq variant (3 × 60).")
note.font = Font(name='Arial', size=10, italic=True)
note.alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=6)
ws.row_dimensions[r + 1].height = 42
widths(ws, {'A': 58, 'B': 22, 'C': 12, 'D': 14, 'E': 16, 'F': 14})

# ---------- Platformaga kiritish ----------
ws = wb.create_sheet('Platformaga kiritish')
lines = [
    ('Saytda kursni shu tartibda yarating', 12, True),
    (f"1. Kurs: «{COURSE}» (fan: Ingliz tili). Tavsifni «Kurs kartochkasi» varag'idan oling.", 10, False),
    ("2. Kurs ichida 6 ta modul (bo'lim) oching: M0 … M5 — nomlari «Modullar» varag'ida.", 10, False),
    ("3. Har bir modulga «Darslar» varag'idagi darslarni tartib raqami bo'yicha qo'shing (jami 36 ta). Dars nomi — «Dars nomi (sayt uchun)» ustuni.", 10, False),
    ("4. Har bir darsga: video (10–15 min) + konspekt (inglizcha kalit so'zlar va o'zbekcha tushuntirish bilan) + dars testi.", 10, False),
    ("5. Dars testini «Avtomatik tuzish» bilan yarating: mavzu — «Savollar bazasidagi mavzu» ustunidagi nom, 10 ta savol, qiyinlik 1–5 aralash.", 10, False),
    ("6. Writing va Speaking darslarida testga qo'shimcha topshiriq qo'shing: yozma ish (80–120 so'z) yoki 1 daqiqalik audio javob — o'qituvchi tekshiradi.", 10, False),
    ("7. Modul oxirida modul testi (30 savol, shu modul mavzularidan). M0 da diagnostik test (40 savol), M5 da 3 ta to'liq mock (60 savol, 70 min + Writing).", 10, False),
    ("8. Savollarni platforma namunasi (savollar-namuna.xlsx) bo'yicha yuklang: «Mavzu» ustuniga «Savollar bazasi» varag'idagi nomni aynan yozing — aks holda test avtomatik tuzilmaydi.", 10, False),
    ('Nega aynan shu tartib', 12, True),
    ("• Avval Present tenses va savol tuzish: kundalik muloqotning poydevori, keyingi barcha zamonlar shularga tayanadi.", 10, False),
    ("• Keyin o'tmish (Past Simple → Past Continuous → Present Perfect): hikoya qilish va tajriba haqida gapirish A2+ darajaning asosiy talabi.", 10, False),
    ("• Kelajak, modal fe'llar va shart gaplar B1 bosqichida: ular oldingi zamonlarni yaxshi bilishni talab qiladi (CEFR / Core Inventory ketma-ketligi).", 10, False),
    ("• Har modulda lug'at, reading/listening va writing/speaking darsi bor — grammatika darhol amaliy ko'nikmada qo'llanadi.", 10, False),
    ("• M5 — takrorlash va 3 ta mock: diagnostika bilan solishtirib, o'sish o'lchanadi.", 10, False),
]
for i, (t, size, bold) in enumerate(lines, start=1):
    c = ws.cell(i, 2, t); c.font = Font(name='Arial', size=size, bold=bold)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[i].height = 30 if len(t) > 110 else 18
widths(ws, {'A': 6, 'B': 100})

wb.save(OUT)
print('saved', OUT, len(topics), 'mavzu')
